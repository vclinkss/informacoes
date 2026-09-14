import os
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

from db import buscar_contatos
from gerar_grafico import gerar_donut
from collections import Counter

PASTA = os.path.dirname(os.path.abspath(__file__))


def wa_link(numero):
    digits = re.sub(r"\D", "", numero)
    if not digits.startswith("55"):
        digits = "55" + digits
    return f"https://wa.me/{digits}"


def main():
    contatos = buscar_contatos()

    # Gera os gráficos atualizados a partir dos dados reais antes de montar a planilha
    caminho_bairro = os.path.join(PASTA, "grafico_bairro.png")
    caminho_escola = os.path.join(PASTA, "grafico_escola.png")
    gerar_donut(Counter(row["bairro"] for row in contatos), caminho_bairro)
    gerar_donut(Counter(row["local_votacao"] for row in contatos), caminho_escola)

    wb = Workbook()
    ws = wb.active
    ws.title = "Contatos"

    FONT_NAME = "Arial"

    ws.merge_cells("A1:H1")
    ws["A1"] = "Relatório de contatos"
    ws["A1"].font = Font(name=FONT_NAME, size=16, bold=True)

    ws.merge_cells("A2:H2")
    ws["A2"] = "Elaborado por Lucas Vinicius"
    ws["A2"].font = Font(name=FONT_NAME, size=10, italic=True, color="666666")

    ws.merge_cells("A3:H3")
    ws["A3"] = f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws["A3"].font = Font(name=FONT_NAME, size=10, color="666666")

    headers = ["Líder", "Nome", "Endereço", "Bairro", "WhatsApp", "Local de votação", "Liguei", "Observação"]
    header_row = 5
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=i, value=h)
        cell.font = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F6E56")
        cell.alignment = Alignment(horizontal="left", vertical="center")

    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    r = header_row
    for row in contatos:
        r += 1
        liguei = row["liguei"]
        values = [
            row["lider"], row["nome"], row["endereco"], row["bairro"],
            row["whatsapp"], row["local_votacao"], "Sim" if liguei else "Não", row["observacao"] or "",
        ]
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = Font(name=FONT_NAME, size=11)
            cell.border = border
            cell.alignment = Alignment(horizontal="left", vertical="center")
            if r % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F2F2F2")
            if c == 5:
                cell.hyperlink = wa_link(value)
                cell.font = Font(name=FONT_NAME, size=11, color="1155CC", underline="single")
            if c == 7:
                cell.font = Font(name=FONT_NAME, size=11, bold=True,
                                  color="1F6E56" if liguei else "993C1D")

    widths = [16, 16, 24, 12, 16, 24, 9, 24]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = f"A{header_row + 1}"

    ws_chart = wb.create_sheet("Gráficos")

    ws_chart["A1"] = "Contatos por bairro"
    ws_chart["A1"].font = Font(name=FONT_NAME, size=12, bold=True)
    img1 = XLImage(caminho_bairro)
    img1.width = 420
    img1.height = 320
    ws_chart.add_image(img1, "A2")

    ws_chart["J1"] = "Contatos por local de votação"
    ws_chart["J1"].font = Font(name=FONT_NAME, size=12, bold=True)
    img2 = XLImage(caminho_escola)
    img2.width = 420
    img2.height = 320
    ws_chart.add_image(img2, "J2")

    destino = os.path.join(PASTA, "relatorio_contatos.xlsx")
    wb.save(destino)
    print("ok:", destino)


if __name__ == "__main__":
    main()
